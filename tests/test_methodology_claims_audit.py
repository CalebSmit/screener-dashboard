"""The methodology page's computational sentences, checked against the code that makes them true.

WHY THIS EXISTS - 2026-10-19 (``research/2026-10-19-methodology-claims-audit.md``).

An audit of every sentence on the generated methodology page that states a computation,
threshold, window, data source, filter, weight, rule or count found sixteen that were false or
stale on the live page, every one of them in prose nobody had registered:

* ROIC's invested capital was described as deducting "cash beyond 2% of revenue"; the engine
  also caps the deduction at half of the cash and floors invested capital at 10% of total
  assets, and the cap bound for **334 of 452** stocks with a ROIC on 2026-10-19;
* the scored dataset was said to be "refreshed daily for prices, weekly for fundamentals"; the
  only cache is the scored dataset, reused only on the day it was written;
* the coverage discount's worked example used 56% coverage, below the coverage filter's bar, and
  the filter itself was said to exclude "fewer than 60%" while it rounds down to whole metrics;
* the missing-data alert was said to cover "any metric"; it read a hand-kept list of 26 that left
  out eight weighted metrics, FY1 EPS Revision (the Revisions category's heaviest) among them;
* the factor-correlation matrix was said to be of the eight category scores; it is of every
  metric's percentile;
* the workbook's dashboard sheet was said to hold the top 50 with all eight category scores and
  both trap flags; it holds the top 25 with three scores and the value-trap flag; the Factor
  Scores sheet was said to carry raw metrics it does not carry; the DataValidation sheet was said
  to show the "portfolio" stocks' market cap, revenue, net income and EPS;
* ``_data_source`` was said to read "yfinance", "cache" or "sample"; it reads ``quarterly`` or
  ``annual``;
* the EPS-basis flag and Limitation 2 still described the forward-EPS construction replaced on
  2026-10-09;
* weight sensitivity was said to perturb each weight "±5%"; it moves it 5 percentage points,
  which for Size's 5% is 0% and 10%;
* the bank set omitted a sub-industry the code lists; the size tilt's 2026-09-03 figures were
  called "the current universe"; the summary listed a liquidity filter, which only the workbook's
  portfolio applies, as a step of the ranking.

Each test reads the generated page (``build_screener_overview``, rule 10), and checks the
**assertion** - against the engine's behaviour, the published payload, or the function the page
and the run now share - not merely that a sentence is present.
"""

from __future__ import annotations

import inspect
import json
import re
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import yaml

import factor_engine as fe
import portfolio_constructor as pc
import run_screener as rs

ROOT = Path(__file__).resolve().parent.parent
PAYLOAD = ROOT / "dashboard_data.js"


@pytest.fixture(scope="module")
def cfg() -> dict:
    return yaml.safe_load((ROOT / "config.yaml").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def page(cfg) -> str:
    return rs.build_screener_overview(cfg)


@pytest.fixture(scope="module")
def payload() -> dict:
    if not PAYLOAD.exists():
        pytest.skip("dashboard_data.js not present")
    text = PAYLOAD.read_text(encoding="utf-8", errors="replace")
    return json.loads(text[text.find("{"):text.rfind("}") + 1])


def _num(x) -> bool:
    return x is not None and not (isinstance(x, float) and np.isnan(x))


# ---------------------------------------------------------------------------
# ROIC's invested capital (claim overview.roic_invested_capital)
# ---------------------------------------------------------------------------

_ROIC_RE = (r"Excess cash is cash beyond ([\d.]+)% of revenue, but never more than ([\d.]+)% "
            r"of the cash, and invested capital is floored at ([\d.]+)% of total assets")


def _roic_from_page(inp: dict, op: float, cap: float | None, floor: float | None):
    """ROIC as the page states it, from a stock's published inputs."""
    ebit = inp.get("ebit")
    if not _num(ebit):
        return None
    tax, pre = inp.get("incomeTaxExpense"), inp.get("pretaxIncome")
    rate = 0.21
    if _num(pre) and pre <= 0:
        rate = 0.0
    elif _num(tax) and _num(pre) and pre > 0:
        rate = max(0.0, min(tax / pre, 0.5))
    eq, debt, cash = inp.get("totalEquity"), inp.get("totalDebt_bs"), inp.get("cash_bs")
    if not (_num(eq) and _num(debt) and _num(cash)):
        return None
    rev, ta = inp.get("totalRevenue"), inp.get("totalAssets")
    excess = max(0.0, cash - (op * rev if _num(rev) and rev > 0 else 0.0))
    if cap is not None:
        excess = min(excess, cap * cash)
    ic = eq + debt - excess
    if floor is not None and _num(ta) and ta > 0:
        ic = max(ic, floor * ta)
    return ebit * (1 - rate) / ic if ic > 0 else None


def _roic_match_rate(payload, op, cap, floor):
    ok = total = 0
    for s in payload["stock_detail"].values():
        pub = s["raw"].get("roic")
        if pub is None:
            continue
        calc = _roic_from_page(s.get("inp") or {}, op, cap, floor)
        if calc is None:
            continue
        total += 1
        ok += abs(calc - pub) <= 1e-4 * max(1.0, abs(calc)) + 1e-4
    return ok, total


def test_the_roic_formula_on_the_page_rebuilds_the_scored_roic(page, payload):
    """Parse the three shares from the page, put every stock's published inputs through them,
    and get the ROIC that was scored. The old wording - the 2% rule alone - must not."""
    m = re.search(_ROIC_RE, page)
    assert m, "the ROIC row no longer states its excess-cash cap and invested-capital floor"
    op, cap, floor = (float(x) / 100 for x in m.groups())

    ok, total = _roic_match_rate(payload, op, cap, floor)
    assert total > 300, f"only {total} stocks carry ROIC and its inputs"
    assert ok / total >= 0.99, f"the page's ROIC formula rebuilds {ok} of {total}"

    old_ok, old_total = _roic_match_rate(payload, op, None, None)
    assert old_ok / old_total < 0.6, (
        "the 2%-of-revenue rule alone now reproduces most stocks' ROIC - the cap no longer "
        "binds widely, so re-measure the claim in research/2026-10-19-methodology-claims-audit.md")


def test_the_page_states_the_engines_own_roic_constants(page):
    m = re.search(_ROIC_RE, page)
    assert m
    assert [float(x) for x in m.groups()] == pytest.approx(
        [fe.ROIC_OPERATING_CASH_SHARE * 100, fe.ROIC_EXCESS_CASH_CAP * 100, fe.ROIC_IC_FLOOR_SHARE * 100])
    src = inspect.getsource(fe.compute_metrics)
    for const in ("ROIC_OPERATING_CASH_SHARE", "ROIC_EXCESS_CASH_CAP", "ROIC_IC_FLOOR_SHARE"):
        assert const in src, f"compute_metrics no longer reads {const}; the page and the engine can drift"
    # The rationale and the design-decisions row say the same, not the bare 2% rule.
    assert "deducts only *excess* cash (cash beyond 2% of revenue)" not in page
    assert re.search(r"ROIC excess cash deduction\*\* \(cash beyond [\d.]+% of revenue, at most [\d.]+% of cash\)", page)


# ---------------------------------------------------------------------------
# The scored-data cache (claim overview.cache_window)
# ---------------------------------------------------------------------------

def test_the_cache_sentence_states_the_window_the_run_applies(page, cfg):
    days = fe.factor_scores_cache_max_age_days(cfg.get("caching", {}))
    if days == 1:
        assert "reused only on the day it was written" in page
    else:
        assert f"for {days} calendar days, counting the day it was written" in page
    today = datetime(2026, 10, 19)
    # Reusable on its own day ...
    assert fe.cache_is_usable(today, days, now=today + timedelta(hours=23))
    # ... and not on the day the window ends, whatever the hour.
    assert not fe.cache_is_usable(today, days, now=today + timedelta(days=days, minutes=1))
    # The run's freshness check reads the same function.
    assert "factor_scores_cache_max_age_days(" in inspect.getsource(rs.run_factor_engine)


def test_there_is_no_separate_weekly_fundamentals_cache(page):
    """The page said fundamentals were cached for a week. Nothing reads
    ``fundamental_data_refresh_days`` except the scored-data bound, which takes the minimum."""
    assert "weekly for fundamentals" not in page
    key = "fundamental_data_refresh_days"
    assert inspect.getsource(fe).count(key) == inspect.getsource(fe.factor_scores_cache_max_age_days).count(key), (
        "factor_engine reads the fundamentals refresh interval somewhere new - if a fundamentals "
        "cache now exists, the page's caching sentence needs rewriting")
    assert key not in inspect.getsource(rs)
    assert fe.factor_scores_cache_max_age_days({"price_data_refresh_days": 1,
                                                "fundamental_data_refresh_days": 7}) == 1


# ---------------------------------------------------------------------------
# Coverage filter and discount (claims overview.coverage_filter, overview.coverage_discount)
# ---------------------------------------------------------------------------

_COV_RE = (r"it needs (\d+) of (\d+) \(([\d.]+)%\), (\d+) of (\d+) if it is bank-like \(([\d.]+)%\), "
           r"and (\d+) of (\d+) for another financial \(([\d.]+)%\)")


def test_the_coverage_filter_bars_are_the_ones_the_run_applies(page, cfg):
    m = re.search(_COV_RE, page)
    assert m, "the coverage-filter bullet no longer states its bars in whole metrics"
    g = [float(x) for x in m.groups()]
    gen, bank = fe.weighted_metric_sets(cfg)
    pct = cfg["data_quality"]["min_data_coverage_pct"]
    for k, n, shown, expect_n in ((g[0], g[1], g[2], len(gen)), (g[3], g[4], g[5], len(bank)),
                                  (g[6], g[7], g[8], len(gen) - 1)):
        assert n == expect_n
        assert k == fe.coverage_minimum(int(n), pct)
        assert shown == pytest.approx(round(k / n * 100, 1))
        # rounded down: the bar sits at or below the configured share, never above it
        assert k / n <= pct / 100 and (k + 1) / n > pct / 100
    # The run's filter reads the same function.
    assert "coverage_minimum(" in inspect.getsource(rs.run_factor_engine)


def test_every_ranked_stock_clears_the_stated_bar(payload, cfg):
    pct = cfg["data_quality"]["min_data_coverage_pct"]
    below = [t for t, s in payload["stock_detail"].items()
             if s.get("cov") and s["cov"]["n"] < fe.coverage_minimum(s["cov"]["of"], pct)]
    assert not below, f"ranked although under the coverage filter's bar: {below[:10]}"


def test_the_discount_example_is_a_coverage_a_ranked_stock_can_have(page, cfg):
    m = re.search(r"a stock with (\d+) of (\d+) metrics, ([\d.]+)% coverage, loses ([\d.]+)% of its "
                  r"composite; the deepest cut the coverage filter leaves possible is ([\d.]+)%, "
                  r"at (\d+) of (\d+)\)", page)
    assert m, "the coverage-discount example has moved"
    k, n, shown, loss, deepest, fk, fn = (float(x) for x in m.groups())
    cd = cfg["data_quality"]["coverage_discount"]
    pct = cfg["data_quality"]["min_data_coverage_pct"]
    assert k >= fe.coverage_minimum(int(n), pct), "the example is a coverage the filter excludes"
    assert loss == pytest.approx((cd["threshold"] - k / n) * cd["penalty_rate"] * 100, abs=0.006)
    gen, bank = fe.weighted_metric_sets(cfg)
    floors = [fe.coverage_minimum(x, pct) / x for x in (len(gen), len(bank), len(gen) - 1)]
    assert fk / fn == pytest.approx(min(floors))
    assert deepest == pytest.approx((cd["threshold"] - min(floors)) * cd["penalty_rate"] * 100, abs=0.006)
    assert "56% coverage" not in page


# ---------------------------------------------------------------------------
# Missing-data alert (claim overview.metric_alerts)
# ---------------------------------------------------------------------------

def test_the_missing_data_alert_covers_every_metric_that_carries_weight(page, cfg):
    gen, bank = fe.weighted_metric_sets(cfg)
    alerted = {m for _, m in rs.drift_alert_metrics(cfg)}
    assert alerted == gen | bank
    assert "fy1_revision_3m" in alerted, "the Revisions category's heaviest metric must be alerted on"
    assert f"missing for more than {cfg['data_quality']['metric_alert_threshold_pct']}% of the stocks it applies to" in page
    src = inspect.getsource(rs.run_factor_engine)
    assert "drift_alert_metrics(cfg)" in src and "labels = _alert_metrics" in src
    # The list is taken before the auto-reduce can zero a weight: a metric reduced for being
    # mostly missing is exactly the one to alert on.
    assert src.index("drift_alert_metrics(cfg)") < src.index("_auto_reduce_high_nan_metrics(df, cfg")


# ---------------------------------------------------------------------------
# Category auto-disable (claim overview.auto_disable)
# ---------------------------------------------------------------------------

def _frame_with_revisions_coverage(cfg, share: float, n: int = 200) -> pd.DataFrame:
    rng = np.random.default_rng(7)
    df = pd.DataFrame({"Ticker": [f"T{i}" for i in range(n)]})
    for m in rs.weighted_category_metrics(cfg, "revisions"):
        v = rng.normal(size=n)
        v[rng.random(n) >= share] = np.nan
        df[m] = v
    return df


def test_a_sparse_category_is_switched_off_and_its_weight_shared_in_proportion(page, cfg):
    thr = rs.CATEGORY_AUTO_DISABLE_PCT
    assert f"If fewer than {thr}% of the values the Revisions or Investment category's weighted metrics" in page
    assert f"When coverage drops below {thr}%" in page

    c = json.loads(json.dumps(cfg))
    disabled, pct = rs.auto_disable_sparse_category(_frame_with_revisions_coverage(c, 0.20), c, "revisions")
    assert disabled and pct < thr
    fw0, fw1 = cfg["factor_weights"], c["factor_weights"]
    assert fw1["revisions"] == 0
    assert sum(fw1.values()) == pytest.approx(100, abs=0.05)
    others = [k for k in fw0 if k != "revisions"]
    for a in others:
        for b in others:
            assert fw1[a] * fw0[b] == pytest.approx(fw1[b] * fw0[a], rel=2e-3)  # in proportion

    c2 = json.loads(json.dumps(cfg))
    disabled2, pct2 = rs.auto_disable_sparse_category(_frame_with_revisions_coverage(c2, 0.45), c2, "revisions")
    assert not disabled2 and pct2 >= thr and c2["factor_weights"] == cfg["factor_weights"]


def _frame_with_exact_coverage(cfg, cat: str, present: int, n: int = 1000) -> pd.DataFrame:
    """Every weighted metric of ``cat`` holds a value for exactly ``present`` of ``n`` stocks."""
    df = pd.DataFrame({"Ticker": [f"T{i}" for i in range(n)]})
    for m in rs.weighted_category_metrics(cfg, cat):
        v = np.arange(n, dtype=float)
        v[present:] = np.nan
        df[m] = v
    return df


@pytest.mark.parametrize("cat", ["revisions", "investment"])
def test_the_auto_disable_bar_is_thirty_percent_and_thirty_itself_is_kept(page, cfg, cat):
    """Claims-mutation audit, 2026-10-26 (``plan/claims-mutation-audit.md`` row 68).

    The test above reads the bar from ``rs.CATEGORY_AUTO_DISABLE_PCT`` and builds both its page
    sentence and its probes (20% and 45%) from it, so ``CATEGORY_AUTO_DISABLE_PCT = 40`` passed:
    the page is generated from the same constant and 20/45 straddle 40 too. The register states
    **30%**, so the figure is written here as a literal and probed on both sides of it - 29.9%
    is switched off, exactly 30.0% is kept ("fewer than 30%") - through the function the run
    calls, with its default threshold, for both categories the page names."""
    assert "If fewer than 30% of the values the Revisions or Investment category's weighted metrics" in page
    c = json.loads(json.dumps(cfg))
    disabled, pct = rs.auto_disable_sparse_category(_frame_with_exact_coverage(c, cat, 299), c, cat)
    assert pct == pytest.approx(29.9) and disabled, "29.9% coverage must switch the category off"
    assert c["factor_weights"][cat] == 0
    c = json.loads(json.dumps(cfg))
    disabled, pct = rs.auto_disable_sparse_category(_frame_with_exact_coverage(c, cat, 300), c, cat)
    assert pct == pytest.approx(30.0) and not disabled, "exactly 30% is not 'fewer than 30%'"
    assert c["factor_weights"] == cfg["factor_weights"]


# ---------------------------------------------------------------------------
# Weight sensitivity (claim overview.weight_sensitivity)
# ---------------------------------------------------------------------------

def test_weight_sensitivity_moves_each_weight_by_percentage_points(page, cfg):
    pts, top = rs.SENSITIVITY_PERTURBATION_PTS, rs.SENSITIVITY_TOP_N
    assert f"up and down by {rs._pts(pts)} percentage points" in page
    assert f"top-{top} ranking" in page and "±5%" not in page
    size = cfg["factor_weights"]["size"]
    assert f"Size's {size}% is tried at {rs._pts(max(size - pts, 0))}% and {rs._pts(size + pts)}%" in page

    rng = np.random.default_rng(3)
    n = 60
    df = pd.DataFrame({"Ticker": [f"T{i}" for i in range(n)]})
    for cat in cfg["factor_weights"]:
        df[f"{cat}_score"] = rng.uniform(0, 100, n)
    df["Rank"] = np.arange(1, n + 1)
    res = fe.run_weight_sensitivity(df, cfg, perturbation_pct=pts, top_n=top)
    for _, r in res.iterrows():
        assert abs(r["perturbed_weight"] - r["original_weight"]) == pytest.approx(pts)
    assert {"+", "-"} <= set(res.loc[res["category"] == "size", "direction"])
    assert "perturbation_pct=SENSITIVITY_PERTURBATION_PTS" in inspect.getsource(rs.run_factor_engine)


def test_weight_sensitivity_compares_the_top_twenty(page, cfg):
    """Claims-mutation audit, 2026-10-26 (row 70). The test above builds its expected page text
    and its call from ``rs.SENSITIVITY_TOP_N``, so ``SENSITIVITY_TOP_N = 25`` passed it. The
    register says the **top 20**; it is a literal here, read off the page and off the set the
    engine actually compares when called the way the run calls it."""
    assert "top-20 ranking" in page and "two top-20 sets" in page
    src = inspect.getsource(rs.run_factor_engine)
    assert "top_n=SENSITIVITY_TOP_N" in src, "the run no longer passes the stated top-N"
    rng = np.random.default_rng(11)
    n = 60
    df = pd.DataFrame({"Ticker": [f"T{i}" for i in range(n)]})
    for cat in cfg["factor_weights"]:
        df[f"{cat}_score"] = rng.uniform(0, 100, n)
    df["Rank"] = np.arange(1, n + 1)
    res = fe.run_weight_sensitivity(df, cfg, perturbation_pct=rs.SENSITIVITY_PERTURBATION_PTS,
                                    top_n=rs.SENSITIVITY_TOP_N)
    # changed is the symmetric difference, so |baseline| + |perturbed| = 2 x unchanged + changed
    assert len(res) and ((2 * res["top_n_unchanged"] + res["top_n_changed"]) == 2 * 20).all()


# ---------------------------------------------------------------------------
# Factor correlation (claim overview.factor_correlation)
# ---------------------------------------------------------------------------

def test_the_correlation_matrix_is_of_metrics_not_categories(page):
    rng = np.random.default_rng(1)
    df = pd.DataFrame({f"{m}_pct": rng.uniform(0, 100, 40) for m in fe.METRIC_COLS})
    for cat in fe.CAT_METRICS:
        df[f"{cat}_score"] = rng.uniform(0, 100, 40)
    corr = fe.compute_factor_correlation(df)
    assert set(corr.columns) == {f"{m}_pct" for m in fe.METRIC_COLS}
    assert not any(c.endswith("_score") for c in corr.columns)
    assert "metric against metric, not category against category" in page
    assert not re.search(r"correlation matrix of all \d+ category scores", page)
    assert "they share five months" in page and "~6 months of overlap" not in page


# ---------------------------------------------------------------------------
# The workbook's sheets (claim overview.workbook_sheets)
# ---------------------------------------------------------------------------

def _scored_frame(n: int = 40) -> pd.DataFrame:
    rng = np.random.default_rng(5)
    df = pd.DataFrame({
        "Ticker": [f"T{i:02d}" for i in range(n)], "Company": [f"Co {i}" for i in range(n)],
        "Sector": ["Industrials", "Energy", "Utilities", "Financials"] * (n // 4),
    })
    for cat in fe.CAT_METRICS:
        df[f"{cat}_score"] = rng.uniform(0, 100, n)
        df[f"{cat}_contrib"] = rng.uniform(0, 20, n)
    for m in fe.METRIC_COLS:
        df[m] = rng.normal(size=n)
    df["Composite"] = np.sort(rng.uniform(20, 80, n))[::-1]
    df["Composite_Pct"] = df["Composite"].rank(pct=True) * 100
    df["Composite_Confidence"] = 80.0
    df["Rank"] = np.arange(1, n + 1)
    for col in ("Value_Trap_Flag", "Growth_Trap_Flag", "Financial_Sector_Caveat", "_is_bank_like",
                "_beneish_flag", "_eps_basis_mismatch", "_stale_data", "_ltm_annualized",
                "_channel_stuffing_flag", "_roic_ic_floored"):
        df[col] = False
    df["Value_Trap_Severity"] = 0.0
    df["Growth_Trap_Severity"] = 0.0
    df["_ev_flag"] = None
    df["_beta_overlap_pct"] = 50.0
    return df


def _headers(ws, row: int) -> list:
    return [c.value for c in ws[row] if c.value is not None]


def test_the_dashboard_sheet_holds_what_the_page_says(page):
    from openpyxl import Workbook
    n = pc.DASHBOARD_TOP_N
    assert (f"The top {n} stocks by composite, formatted for quick review: rank, sector, composite "
            "(quartile-colored), the Valuation, Quality and Momentum scores and the value-trap flag") in page
    assert "The top 50 stocks, formatted for quick review" not in page
    wb = Workbook()
    pc.write_screener_dashboard(wb, _scored_frame())
    ws = wb["ScreenerDashboard"]
    heads = _headers(ws, 5)
    assert heads == list(pc.DASHBOARD_TOP_COLUMNS)
    scores = [h for h in heads if h.endswith("_Pct")]
    assert scores == ["Val_Pct", "Qual_Pct", "Mom_Pct"]
    assert "Growth_Trap_Flag" not in heads and not any("Severity" in h for h in heads)
    tickers = [ws.cell(row=r, column=2).value for r in range(6, 6 + n)]
    assert tickers == [f"T{i:02d}" for i in range(n)]


def test_the_factor_scores_sheet_carries_no_raw_metrics(page):
    from openpyxl import Workbook
    assert "Raw metric values are not on this sheet" in page
    wb = Workbook()
    pc.write_disclaimer_sheet(wb)
    pc.write_factor_scores_sheet(wb, _scored_frame())
    heads = _headers(wb["FactorScores"], 1)
    for h in ("Composite", "Composite_Pct", "Confidence", "Rank", "VT_Severity", "GT_Severity",
              "Value_Trap_Flag", "Growth_Trap_Flag", "Beneish_Flag", "Fin_Caveat", "Is_Bank"):
        assert h in heads, h
    assert sum(h.endswith("_Contrib") for h in heads) == 8
    assert sum(h.endswith("_Pct") and h != "Composite_Pct" for h in heads) == 8
    assert not set(heads) & set(fe.METRIC_COLS), "raw metrics are on the sheet after all"


def test_the_validation_sheet_is_the_top_of_the_ranking_with_seven_flags_highlighted(page):
    from openpyxl import Workbook
    n = pc.DATA_VALIDATION_TOP_N
    assert f"The top {n} stocks of the ranking (not the portfolio's holdings)" in page
    assert "highlights seven types of potential issue" in page
    df = _scored_frame()
    df.loc[0, "Value_Trap_Flag"] = True   # the top stock is trap-flagged: no portfolio would hold it
    flags = ["_eps_basis_mismatch", "_stale_data", "_ltm_annualized", "_channel_stuffing_flag",
             "_beneish_flag", "_roic_ic_floored"]
    for i, f in enumerate(flags, 1):
        df.loc[i, f] = True
    df.loc[len(flags) + 1, "_ev_flag"] = "API EV=10B vs computed=5B"
    wb = Workbook()
    pc.write_data_validation_sheet(wb, df, top_n=n)
    ws = wb["DataValidation"]
    heads = _headers(ws, 3)
    tick_col = heads.index("Ticker") + 1
    shown = [ws.cell(row=r, column=tick_col).value for r in range(4, 4 + n)]
    assert shown == [f"T{i:02d}" for i in range(n)], "not the top of the ranking"
    for label in ("EV/EBITDA (raw)", "FCF Yield (raw)", "Earnings Yield (raw)", "ROIC (raw)", "D/E (raw)",
                  "Fwd EPS Growth (raw)", "Rev Growth (raw)", "12-1M Return (raw)", "Volatility (raw)",
                  "Beta (raw)", "Analyst Surprise (raw)", "Short Interest (raw)", "Price", "Analyst Target",
                  "# Analysts", "Metrics Available", "Beta Overlap %"):
        assert label in heads, label
    assert not {"Market Cap", "Revenue", "Net Income", "EPS"} & set(heads)

    def filled(row, header):
        return ws.cell(row=row, column=heads.index(header) + 1).fill.fgColor.rgb not in (None, "00000000")

    for i, h in enumerate(["EPS Mismatch?", "Stale Data?", "LTM Partial?", "Channel Stuffing?",
                           "Beneish Flag", "ROIC IC-Floor?"], 1):
        assert filled(4 + i, h), f"{h} is not highlighted"
    assert filled(4 + len(flags) + 1, "EV Discrepancy")
    assert not any(filled(r, "Beta Overlap %") for r in range(4, 4 + n)), "beta overlap is shown, not highlighted"


def test_the_workbook_the_run_writes_holds_the_top_25_and_validates_the_top_10(page, tmp_path, monkeypatch):
    """Claims-mutation audit, 2026-10-26 (rows 66-67). The two tests above take the row counts
    from ``pc.DASHBOARD_TOP_N`` and ``pc.DATA_VALIDATION_TOP_N`` and the second passes its own
    ``top_n`` to the sheet writer, so ``DASHBOARD_TOP_N = 20`` and ``DATA_VALIDATION_TOP_N = 12``
    both passed every registered check while the page and the workbook moved together. The
    register states **25** and **10**: they are literals here, both on the page and in the
    workbook ``write_full_excel`` - the function the run calls - actually saves, with the row
    after each block checked empty so a longer list fails as well as a shorter one."""
    from openpyxl import load_workbook
    assert "The top 25 stocks by composite, formatted for quick review" in page
    assert "The top 10 stocks of the ranking (not the portfolio's holdings)" in page
    # The portfolio sheet is not part of this claim and needs a portfolio; skip it.
    monkeypatch.setattr(pc, "write_model_portfolio_sheet", lambda *a, **k: None)
    out = tmp_path / "workbook.xlsx"
    path = pc.write_full_excel(_scored_frame(), pd.DataFrame(), {}, {"output": {"excel_file": str(out)}})
    assert Path(path) == out
    wb = load_workbook(out)
    ws = wb["ScreenerDashboard"]
    assert [ws.cell(row=r, column=2).value for r in range(6, 31)] == [f"T{i:02d}" for i in range(25)]
    assert ws.cell(row=31, column=2).value is None, "more than 25 rows on the dashboard sheet"
    ws = wb["DataValidation"]
    tick_col = _headers(ws, 3).index("Ticker") + 1
    assert [ws.cell(row=r, column=tick_col).value for r in range(4, 14)] == [f"T{i:02d}" for i in range(10)]
    assert ws.cell(row=14, column=tick_col).value is None, "more than 10 rows on the validation sheet"


# ---------------------------------------------------------------------------
# EPS basis flag and Limitation 2 (claim overview.eps_basis)
# ---------------------------------------------------------------------------

def test_the_eps_flag_no_longer_removes_the_scored_growth_figure(page, payload):
    assert "Since 2026-10-09 the scored Forward EPS Growth does not use that pair" in page
    assert "Since 2026-10-09 Forward EPS Growth avoids the mix" in page
    flagged = [s for s in payload["stock_detail"].values() if s.get("eps_mismatch")]
    assert flagged, "no stock carries the EPS-basis flag; re-measure before relying on this claim"
    twelve_month = [s for s in flagged if (s.get("inp") or {}).get("feg_f12") is not None]
    # "for most flagged stocks": the twelve-month construction, where the flag does not apply
    assert len(twelve_month) > len(flagged) / 2
    kept = [s for s in twelve_month if s["raw"].get("forward_eps_growth") is not None
            or (s["inp"].get("feg_b12") is not None and s["inp"]["feg_b12"] <= 0)]
    assert len(kept) == len(twelve_month), "a flagged stock on the 12-month basis lost its growth figure"
    # The fallback keeps the ratio rule: forward over trailing outside [0.3, 2] has no growth figure.
    for s in payload["stock_detail"].values():
        inp = s.get("inp") or {}
        if inp.get("feg_f12") is not None:
            continue
        f, t = inp.get("forwardEps"), inp.get("trailingEps")
        if _num(f) and _num(t) and abs(t) > 0.01 and not 0.3 <= f / t <= 2.0:
            assert s["raw"].get("forward_eps_growth") is None


def _feg_row(fwd: float, trail: float, twelve_month: bool) -> dict:
    """One stock through ``compute_metrics``; ``twelve_month`` supplies the consensus inputs the
    2026-10-09 construction needs (so the fallback is not used)."""
    from datetime import timezone
    from tests.test_metrics import _compute_one, _make_rec
    extra = {}
    if twelve_month:
        # clock: a consensus year-end six months out, which is mid-range on any day of the week
        # and far from the construction's 0-to-12-month boundaries.
        end = datetime.now(timezone.utc) + timedelta(days=180)
        extra = {"_fy1_eps_current": 5.0, "_fy2_eps_current": 6.0, "_next_fy_end": int(end.timestamp()),
                 "_eps_quarters": json.dumps([[f"2026-0{i + 1}-30", 1.0, 1.0, 0.0] for i in range(4)])}
    return _compute_one(_make_rec(forwardEps=fwd, trailingEps=trail, **extra))


def test_the_engine_flags_outside_03_to_2x_and_drops_growth_only_on_the_fallback():
    """Claims-mutation audit, 2026-10-26 (rows 71-73). The test above reads the committed
    payload, so a change to ``compute_metrics`` could not fail it before the data loop's next
    run: widening the fallback's ratio rule to 2.5x, the flag's bar to 2.5x, or applying the
    rule to the twelve-month construction all passed gate 1. This runs the engine, with the
    stated 0.3x and 2x as literals probed on both sides."""
    trail = 5.0
    # The flag: forward over GAAP trailing outside 0.3-2x, either side of each bar.
    for fwd, flagged in ((10.1, True), (9.9, False), (1.45, True), (1.55, False)):
        r = _feg_row(fwd, trail, twelve_month=False)
        assert bool(r["_eps_basis_mismatch"]) is flagged, (fwd / trail, r["_eps_basis_mismatch"])
        # The fallback construction still sets growth to NaN on an extreme ratio, and only then.
        assert r["_feg_basis"] == "fy2_over_trailing"
        assert bool(np.isnan(r["forward_eps_growth"])) is flagged, (fwd / trail, r["forward_eps_growth"])
    # The twelve-month construction does not use that pair, so the flag does not remove it.
    for fwd in (10.1, 1.45):
        r = _feg_row(fwd, trail, twelve_month=True)
        assert r["_eps_basis_mismatch"] and r["_feg_basis"] == "msci_12m"
        assert np.isfinite(r["forward_eps_growth"]), "the flag removed a twelve-month growth figure"


# ---------------------------------------------------------------------------
# Provenance and bank-set wording (claims overview.data_source, overview.bank_set)
# ---------------------------------------------------------------------------

def test_data_source_names_the_values_the_fetch_writes(page, payload):
    written = set(re.findall(r'rec\["_data_source"\] = "(\w+)"', inspect.getsource(fe)))
    assert written == {"annual", "quarterly"}, written
    for v in written:
        assert f"`{v}`" in page
    assert '"yfinance", "cache", "sample"' not in page
    published = {s.get("data_source") for s in payload["stock_detail"].values()}
    assert published <= written | {None, "unknown"}, published


_SUBINDUSTRY_ON_PAGE = {
    "Diversified Banks": "Diversified and Regional Banks", "Regional Banks": "Diversified and Regional Banks",
    "Consumer Finance": "Consumer Finance", "Commercial & Residential Mortgage Finance": "Mortgage Finance",
    "Life & Health Insurance": "Life & Health", "Multi-line Insurance": "Multi-line",
    "Property & Casualty Insurance": "Property & Casualty Insurance", "Reinsurance": "Reinsurance",
    "Multi-Sector Holdings": "Multi-Sector Holdings", "Investment Banking & Brokerage": "Investment Banking & Brokerage",
    "Diversified Capital Markets": "Diversified Capital Markets",
    "Insurance Brokers": "Insurance Brokers", "Asset Management & Custody Banks": "Asset Management",
    "Financial Exchanges & Data": "Financial Exchanges & Data",
    "Transaction & Payment Processing Services": "Payment Processing",
}


def test_every_sub_industry_the_code_lists_is_named_on_the_page(page):
    section = page.split("## Bank-specific scoring")[1].split("## How the score is calculated")[0]
    bank_line = next(ln for ln in section.splitlines() if ln.startswith("1. **Bank set**"))
    generic_line = next(ln for ln in section.splitlines() if ln.startswith("2. **Generic set**"))
    for sub in fe._BANK_SUBINDUSTRIES:
        assert sub in _SUBINDUSTRY_ON_PAGE, f"new bank sub-industry {sub!r}: name it on the page"
        assert _SUBINDUSTRY_ON_PAGE[sub] in bank_line, sub
    for sub in fe._GENERIC_FIN_SUBINDUSTRIES:
        assert _SUBINDUSTRY_ON_PAGE[sub] in generic_line, sub
    assert fe._is_bank_like("X", "Financials", "", "Diversified Capital Markets")


# ---------------------------------------------------------------------------
# Counts, the summary, and the dated size figures
# ---------------------------------------------------------------------------

def test_the_overview_states_how_many_metrics_each_kind_of_stock_is_scored_on(page, cfg):
    gen, bank = fe.weighted_metric_sets(cfg)
    assert (f"so a bank-like stock is scored on {len(bank)}, and a financial company outside the bank "
            f"set on {len(gen) - 1}") in page
    assert "each stock is scored on about" not in page


def test_the_liquidity_filter_is_the_portfolios_not_the_rankings(page):
    assert "6. For the model portfolio in the Excel workbook only, apply a liquidity filter" in page
    assert "min_avg_dollar_volume" not in inspect.getsource(rs.run_factor_engine)
    assert "min_avg_dollar_volume" in inspect.getsource(pc.construct_portfolio)


def test_the_size_example_is_dated_not_called_current(page):
    size = page.split("### 7. Size")[1].split("### 8.")[0]
    assert "on the 2026-09-03 universe CAT ($364B) scored 1 and UAL ($35B) scored 59" in size
    assert "on the current universe" not in size


# ---------------------------------------------------------------------------
# The weights, the percentile rule and the trap rules (claims overview.metric_weight_tables,
# overview.percentile_rule, overview.trap_rules)
# ---------------------------------------------------------------------------

_LABEL_TO_METRIC = {v: k for k, v in rs._METRIC_LABELS.items()}


def _tables(block: str) -> list[dict]:
    """Every metric table in ``block``, as {metric: weight}."""
    out, cur = [], None
    for ln in block.splitlines():
        m = re.match(r"^\| \*\*(.+?)\*\* \| ([\d.]+)% \|", ln)
        if m:
            if cur is None:
                cur = {}
            cur[_LABEL_TO_METRIC[m.group(1)]] = float(m.group(2))
        elif cur is not None and not ln.startswith("|"):
            out.append(cur)
            cur = None
    if cur:
        out.append(cur)
    return out


def test_every_weight_table_on_the_page_is_the_engines(page, cfg):
    """Every table, in every category, against ``metric_weight_profiles`` - the tables the stocks
    are scored with - and every heading and the composite line against ``factor_weights``."""
    heads = list(re.finditer(r"^### (\d)\. (.+?) \((\d+)% of final score\)", page, re.M))
    assert len(heads) == 8
    names = {"Valuation": "valuation", "Quality": "quality", "Growth": "growth", "Momentum": "momentum",
             "Risk": "risk", "Analyst Revisions": "revisions", "Size": "size", "Investment": "investment"}
    for i, h in enumerate(heads):
        cat = names[h.group(2)]
        assert float(h.group(3)) == cfg["factor_weights"][cat]
        end = heads[i + 1].start() if i + 1 < len(heads) else page.index("## Bank-specific scoring")
        tables = _tables(page[h.end():end])
        prof = fe.metric_weight_profiles(cfg, cat)
        want = [{m: w * 100 for m, w in prof["generic"].items() if w > 0}]
        if "bank" in prof and cat in ("valuation", "quality"):
            want.append({m: w * 100 for m, w in prof["bank"].items() if w > 0})
        assert len(tables) == len(want), f"{cat}: {len(tables)} tables on the page"
        for got, exp in zip(tables, want):
            assert got.keys() == exp.keys(), cat
            for m in exp:
                assert got[m] == pytest.approx(exp[m]), (cat, m)
    formula = re.search(r"Raw Composite = (.+)", page).group(1)
    for w, label in re.findall(r"([\d.]+)% × (\w+)", formula):
        assert float(w) == cfg["factor_weights"][label.lower()]


def test_the_percentile_rule_is_the_midpoint_rank_with_a_small_sector_fallback(page):
    m = re.search(r"the k-th lowest of n stocks scores \(k − 0\.5\) / n x 100", page)
    assert m, "the midpoint rule sentence has moved"
    s = pd.Series([3.0, 1.0, 2.0, 5.0, 4.0])
    up = fe._directed_pct(s, True)
    assert list(up) == pytest.approx([(k - 0.5) / 5 * 100 for k in (3, 1, 2, 5, 4)])
    assert list(fe._directed_pct(s, False)) == pytest.approx(list(100 - up))
    assert up.mean() == pytest.approx(50.0)
    k = int(re.search(r"When a sector has fewer than (\d+) stocks with valid data", page).group(1))
    assert k == fe.SECTOR_MIN_PEERS
    # Behaviour: a sector with k-1 valued stocks is ranked against the universe, not itself.
    vals = list(range(1, 31))
    df = pd.DataFrame({"Ticker": [f"T{i}" for i in vals],
                       "Sector": ["Small"] * (k - 1) + ["Big"] * (30 - k + 1)})
    for c in fe.METRIC_COLS:
        df[c] = np.nan
    df["roic"] = [float(v) for v in vals]
    out = fe.compute_sector_percentiles(df.copy())
    small = out[out["Sector"] == "Small"]["roic_pct"]
    assert list(small) == pytest.approx([(v - 0.5) / 30 * 100 for v in vals[:k - 1]])


def test_a_sector_with_exactly_the_minimum_is_ranked_within_itself(page):
    """Claim ``overview.percentile_rule``: the universe fallback is for *fewer than* k.

    The test above checks only the k-1 side, so moving the boundary (``<`` to ``<=``:
    a sector with exactly k values also falls back) passed (mutation audit 2026-10-20,
    ``plan/claims-mutation-audit.md``). Here a sector with exactly k valued stocks must
    be ranked against itself - midpoint ranks 1..k over k - and one with k-1 against
    the universe, side by side in the same frame."""
    k = int(re.search(r"When a sector has fewer than (\d+) stocks with valid data", page).group(1))
    n_big = 40
    vals = list(range(1, k + (k - 1) + n_big + 1))
    sectors = ["AtMin"] * k + ["Below"] * (k - 1) + ["Big"] * n_big
    df = pd.DataFrame({"Ticker": [f"T{i}" for i in vals], "Sector": sectors})
    for c in fe.METRIC_COLS:
        df[c] = np.nan
    df["roic"] = [float(v) for v in vals]
    out = fe.compute_sector_percentiles(df.copy())
    at_min = out[out["Sector"] == "AtMin"]["roic_pct"]
    assert list(at_min) == pytest.approx([(j - 0.5) / k * 100 for j in range(1, k + 1)])
    below = out[out["Sector"] == "Below"]["roic_pct"]
    assert list(below) == pytest.approx([(v - 0.5) / len(vals) * 100 for v in vals[k:2 * k - 1]])


def test_the_trap_rules_on_the_page_are_the_rules_the_engine_applies(page, cfg):
    vt = re.search(r"a Valuation score in the top (\d+)% of the universe – \*\*and\*\* it falls in the "
                   r"bottom (\d+)% of \*\*at least two\*\*", page)
    gt = re.search(r"Growth Score is \*\*above\*\* the (\d+)th percentile.*?Quality Score \*\*below\*\* the "
                   r"(\d+)th percentile.*?Revisions Score \*\*below\*\* the (\d+)th percentile", page, re.S)
    assert vt and gt
    top_val, floor = (int(x) / 100 for x in vt.groups())
    g_ceil, g_q, g_r = (int(x) / 100 for x in gt.groups())
    rng = np.random.default_rng(11)
    n = 400
    df = pd.DataFrame({c: rng.uniform(0, 100, n) for c in
                       ("valuation_score", "quality_score", "momentum_score", "revisions_score", "growth_score")})
    out = fe.apply_growth_trap_flags(fe.apply_value_trap_flags(df.copy(), cfg), cfg)
    q = df.quantile
    cheap = df["valuation_score"] >= q(1 - top_val)["valuation_score"]
    weak = sum((df[c] <= q(floor)[c]).astype(int) for c in ("quality_score", "momentum_score", "revisions_score"))
    assert (out["Value_Trap_Flag"] == (cheap & (weak >= 2))).all()
    grow = df["growth_score"] >= q(g_ceil)["growth_score"]
    gt_expect = grow & ((df["quality_score"] <= q(g_q)["quality_score"])
                        | (df["revisions_score"] <= q(g_r)["revisions_score"]))
    assert (out["Growth_Trap_Flag"] == gt_expect).all()


# ---------------------------------------------------------------------------
# Clamps and denominator floors (claim overview.clamps_and_floors)
# ---------------------------------------------------------------------------

def test_the_fetch_applies_the_surprise_floor_the_page_states(page, monkeypatch):
    """Claim ``overview.clamps_and_floors``: 'analyst surprise divides by the estimate
    floored at $0.10' - checked against the code that computes it.

    The test below compares the page with the *committed payload*, which the code under
    test did not produce: halving the floor inside the fetch passed every check the
    code loop runs, and would first have failed in the data loop's publish gate, the
    morning after it merged (mutation audit 2026-10-20, ``plan/claims-mutation-audit.md``).
    Here the fetch runs on a stub with two quarters whose estimates sit below the floor."""
    import yfinance

    floor = float(re.search(r"Analyst surprise uses a \$([\d.]+) floor on estimated EPS", page).group(1))
    # clock: four quarter dates in descending order, all in the past. The fetch takes them as
    # given - no weekday, session or market calendar enters the floor it is being asked about.
    now = pd.Timestamp.now().normalize()
    quarters = [now - pd.Timedelta(days=d) for d in (300, 210, 120, 30)]
    est = [1.00, 2.00, 0.05, 0.04]
    act = [1.05, 2.10, 0.06, 0.08]

    class _Stub:
        info = {"symbol": "SURP", "marketCap": 1.0e10, "currentPrice": 50.0}
        earnings_history = pd.DataFrame({"epsActual": act, "epsEstimate": est}, index=quarters)

        def __getattr__(self, name):
            return None

    monkeypatch.setattr(yfinance, "Ticker", lambda *a, **k: _Stub())
    rec = fe._fetch_single_ticker_inner("SURP")
    want = [(a - e) / max(abs(e), floor) for a, e in zip(act, est)]
    published = [q[3] for q in json.loads(rec["_eps_quarters"])]
    assert published == pytest.approx(want)
    assert rec["analyst_surprise"] == pytest.approx(float(np.median(want)))
    assert want[-1] == pytest.approx(0.04 / floor)      # the floor binds on this quarter


def test_the_engine_clamps_price_target_upside_where_the_page_says(page, cfg):
    """Claim ``overview.clamps_and_floors``: price target upside 'is clamped to the
    configured [-50%, +100%]' - checked on ``compute_metrics`` run with the config.

    The claim's own test reads the committed payload, so it could not see a change to the
    clamp in the code; ``tests/test_metrics.py`` caught one, but it is not in the register
    (it is a class method, which ``checked_by`` cannot name) and runs without the config
    the page quotes (mutation audit 2026-10-20, ``plan/claims-mutation-audit.md``)."""
    m = re.search(r"^\| \*\*Price Target Upside\*\* \|.*?Clamped to \[(-?\d+)%, \+(\d+)%\]", page, re.M)
    lo, hi = int(m.group(1)) / 100, int(m.group(2)) / 100
    base = {"Ticker": "PTU", "sector": "Industrials", "marketCap": 1.0e10,
            "currentPrice": 100.0, "numberOfAnalystOpinions": 5}

    def ptu(target, c=cfg):
        rec = dict(base, targetMeanPrice=target)
        return fe.compute_metrics([rec], pd.Series(dtype=float), cfg=c).iloc[0]["price_target_upside"]

    # "the configured": the clamp follows config, not a constant that happens to agree with it
    other = json.loads(json.dumps(cfg))
    other["metric_clamps"]["price_target_upside"] = [-0.4, 0.8]
    assert ptu(300.0, other) == pytest.approx(0.8) and ptu(10.0, other) == pytest.approx(-0.4)

    assert ptu(100.0 * (1 + hi) * 1.2) == pytest.approx(hi)
    assert ptu(100.0 * (1 + hi) * 0.95) == pytest.approx((1 + hi) * 0.95 - 1)
    assert ptu(100.0 * (1 + lo) * 0.8) == pytest.approx(lo)
    assert ptu(100.0 * (1 + lo) * 1.05) == pytest.approx((1 + lo) * 1.05 - 1)


def test_the_clamps_and_floors_on_the_page_are_the_ones_applied(page, cfg, payload):
    """The metric rows hard-code their clamps; the Data quality bullet reads config. Both must
    be config's, and the published values must be what the stated floors and clamps give."""
    clamps = cfg["metric_clamps"]
    for lab, key in (("Forward EPS Growth", "forward_eps_growth"), ("Price Target Upside", "price_target_upside")):
        m = re.search(rf"^\| \*\*{re.escape(lab)}\*\* \|.*?Clamped to \[(-?\d+)%, \+(\d+)%\]", page, re.M)
        assert m, f"the {lab} row no longer states its clamp"
        assert (int(m.group(1)), int(m.group(2))) == tuple(round(x * 100) for x in clamps[key])
    floors = re.search(r"Analyst surprise uses a \$([\d.]+) floor on estimated EPS; forward EPS growth uses a "
                       r"\$([\d.]+) floor on trailing EPS", page)
    assert floors
    sur_floor, eps_floor = (float(x) for x in floors.groups())
    lo, hi = clamps["forward_eps_growth"]
    n_sur = n_feg = 0
    for s in payload["stock_detail"].values():
        raw, inp = s["raw"], s.get("inp") or {}
        eq4 = s.get("eq4") or []
        q = [r for r in eq4 if r[1] is not None and r[2] is not None and abs(r[2]) > 0.001]
        if raw.get("analyst_surprise") is not None and len(q) >= 2 and len(eq4) == len(q):
            want = float(np.median([(r[1] - r[2]) / max(abs(r[2]), sur_floor) for r in q]))
            assert raw["analyst_surprise"] == pytest.approx(want, abs=1e-6)
            n_sur += 1
        f12, b12 = inp.get("feg_f12"), inp.get("feg_b12")
        if raw.get("forward_eps_growth") is not None and f12 is not None and b12 is not None and b12 > 0:
            want = float(np.clip((f12 - b12) / max(abs(b12), eps_floor), lo, hi))
            assert raw["forward_eps_growth"] == pytest.approx(want, abs=1e-4)
            n_feg += 1
    assert n_sur > 300 and n_feg > 300, (n_sur, n_feg)
